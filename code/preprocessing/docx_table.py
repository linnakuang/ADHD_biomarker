from __future__ import annotations
from pathlib import Path

def _cell_border(cell, top=None, bottom=None):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tcPr.append(borders)
    for side, spec in (("top", top), ("bottom", bottom)):
        if spec is None:
            continue
        ex = borders.find(qn(f"w:{side}"))
        if ex is not None:
            borders.remove(ex)
        size, color = spec
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size)); el.set(qn("w:color"), color)
        borders.append(el)

def _fixed_layout(table):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tblPr = table._element.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr"); table._element.insert(0, tblPr)
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout"); tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

def _render(path, title, col_labels, rows, notes=None, landscape=False,
            col_widths=None, font_pt=8, header_pt=9):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    doc = Document()
    if landscape:
        from docx.enum.section import WD_ORIENT
        sec = doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    heading = doc.add_heading(title, level=2)
    for r in heading.runs:
        r.font.color.rgb = RGBColor(0, 0, 0)

    tbl = doc.add_table(rows=len(rows) + 1, cols=len(col_labels))
    tbl.autofit = False; tbl.allow_autofit = False
    _fixed_layout(tbl)
    if col_widths:
        for row in tbl.rows:
            for j, cell in enumerate(row.cells):
                cell.width = Inches(col_widths[j])

    def _put(cell, val, bold, size):
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(0); para.paragraph_format.space_after = Pt(0)
        run = para.add_run("" if val is None else str(val))
        run.bold = bold; run.font.size = Pt(size)

    for j, lab in enumerate(col_labels):
        cell = tbl.cell(0, j)
        _put(cell, lab, True, header_pt)
        _cell_border(cell, bottom=(12, "000000"))

    for i, (kind, cells) in enumerate(rows, start=1):
        for j, val in enumerate(cells):
            cell = tbl.cell(i, j)
            if kind == "section":
                _put(cell, val, True, header_pt + 1)
                _cell_border(cell, top=(8, "000000"), bottom=(8, "000000"))
            elif kind == "variable":
                _put(cell, val, True, font_pt)
                _cell_border(cell, top=(4, "000000"))
            else:
                _put(cell, val, False, font_pt)

    for j in range(len(col_labels)):
        _cell_border(tbl.cell(len(rows), j), bottom=(12, "000000"))

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path

def write_docx_table(path, title, columns, rows, notes=None, font_pt=7):
    styled = [("data", list(r)) for r in rows]
    return _render(path, title, columns, styled, notes, font_pt=font_pt, header_pt=font_pt + 1)

def write_styled_table(path, title, columns, rows, notes=None, landscape=True, col_widths=None):
    return _render(path, title, columns, rows, notes, landscape=landscape, col_widths=col_widths)
